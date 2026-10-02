#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pembuat belangko penelitian KOSONG (docx dan xlsx).

Belangko hanya berisi format, label kolom, petunjuk pengisian, bahan yang dibutuhkan
untuk analisis, dan langkah analisis umum. Tidak ada isi substantif, tidak ada rumus
perhitungan, tidak ada data. Peneliti mengisi sendiri.

Pemakaian
  python buat_belangko.py --daftar
  python buat_belangko.py --jenis transkrip-wawancara --keluar ./keluaran
  python buat_belangko.py --jenis tabulasi-angket --butir 25 --responden 40 --keluar ./keluaran
  python buat_belangko.py --semua --keluar ./keluaran
  python buat_belangko.py --uji --keluar ./keluaran      # periksa bahwa belangko benar-benar kosong
  python buat_belangko.py --katalog-md                    # cetak katalog dalam markdown
"""
import argparse, os, sys, re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from katalog import FORMS, XLSX_KOLOM, ETIKA_DATA

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as e:
    sys.exit("Pustaka belum terpasang. Jalankan: pip install --break-system-packages python-docx openpyxl\n" + str(e))

FONT = "Times New Roman"
CATATAN_AKHIR = "Belangko kosong. Seluruh isian diisi sendiri oleh peneliti."


# ---------- util docx ----------
def _font(run, size=11, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _para(doc, text="", size=11, bold=False, italic=False, align=None, space_after=4):
    p = doc.add_paragraph()
    if text:
        _font(p.add_run(text), size, bold, italic)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def _shade(cell, hex_color="D9D9D9"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cell_text(cell, text, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if text:
        _font(p.add_run(text), size, bold)
    p.paragraph_format.space_after = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _footer_page_number(section, note):
    p = section.footer.paragraphs[0]
    p.text = ""
    _font(p.add_run(note + "  Halaman "), 9, italic=True)
    run = p.add_run()
    _font(run, 9, italic=True)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); run._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; run._r.append(i)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _new_doc(landscape):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.0))
    _footer_page_number(sec, CATATAN_AKHIR)
    return doc


def _identitas_table(doc, labels):
    if not labels:
        return
    t = doc.add_table(rows=len(labels), cols=2)
    t.style = "Table Grid"
    for i, lab in enumerate(labels):
        _cell_text(t.cell(i, 0), lab, bold=True)
        _cell_text(t.cell(i, 1), "")
        _shade(t.cell(i, 0), "F2F2F2")
        t.rows[i].height = Cm(0.8)
        t.rows[i].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        t.cell(i, 0).width = Cm(7.0)
    _para(doc, "", space_after=6)


def _main_table(doc, kolom, baris, nomor=True, bagian=None, total_cm=25.7):
    if bagian is not None:
        t = doc.add_table(rows=len(bagian) + 1, cols=2)
        t.style = "Table Grid"
        for j, k in enumerate(kolom[:2]):
            _cell_text(t.cell(0, j), k, bold=True, center=True); _shade(t.cell(0, j))
        _repeat_header(t.rows[0])
        for i, b in enumerate(bagian, start=1):
            _cell_text(t.cell(i, 0), b, bold=True)
            _cell_text(t.cell(i, 1), "")
            t.rows[i].height = Cm(1.6)
            t.rows[i].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        return t
    t = doc.add_table(rows=baris + 1, cols=len(kolom))
    t.style = "Table Grid"
    for j, k in enumerate(kolom):
        _cell_text(t.cell(0, j), k, bold=True, center=True)
        _shade(t.cell(0, j))
    _repeat_header(t.rows[0])
    first_is_no = kolom[0].lower().startswith("no")
    for i in range(1, baris + 1):
        for j in range(len(kolom)):
            _cell_text(t.cell(i, j), str(i) if (j == 0 and first_is_no and nomor) else "", center=(j == 0 and first_is_no))
        t.rows[i].height = Cm(1.0)
        t.rows[i].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    _atur_lebar(t, kolom, total_cm)
    return t


def _atur_lebar(t, kolom, total_cm=25.7):
    """Bobot lebar kolom. Kolom nomor sempit, kolom isi lebar."""
    lebar_total = Cm(total_cm)
    pola_lebar = ("isi", "pernyataan", "kutipan", "deskripsi", "catatan", "pertanyaan", "definisi", "uraian", "keputusan", "masukan", "apa ", "mengapa", "alasan", "dampak", "butir", "temuan")
    bobot = []
    for k in kolom:
        kl = k.lower()
        if kl.startswith("no") and len(kl) <= 9:
            bobot.append(0.6)
        elif kl.isdigit() or kl in ("muncul", "paraf", "tanggal", "waktu", "tahun", "kode", "skor penilaian"):
            bobot.append(0.8)
        elif any(w in kl for w in pola_lebar):
            bobot.append(2.2)
        else:
            bobot.append(1.4)
    jml = sum(bobot)
    t.autofit = False
    lebar = [int(lebar_total * b / jml) for b in bobot]
    for j, col in enumerate(t.columns):
        col.width = lebar[j]
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = lebar[j]


def _bullets(doc, items, numbered=False):
    for n, it in enumerate(items, start=1):
        p = doc.add_paragraph()
        _font(p.add_run((f"{n}. " if numbered else "- ") + it), 11)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)


def _panduan_pages(doc, f):
    doc.add_page_break()
    _para(doc, "Petunjuk Pengisian", 13, bold=True, space_after=6)
    _bullets(doc, f["petunjuk"] + [ETIKA_DATA], numbered=True)
    _para(doc, "", space_after=4)
    _para(doc, "Bahan yang Dibutuhkan untuk Analisis", 13, bold=True, space_after=6)
    _bullets(doc, f["bahan"])
    _para(doc, "", space_after=4)
    _para(doc, "Langkah Analisis Umum", 13, bold=True, space_after=6)
    _bullets(doc, f["langkah"], numbered=True)
    _para(doc, "", space_after=4)
    _para(doc, "Catatan. Langkah di atas bersifat umum dan tidak bergantung aplikasi. Untuk tutorial menurut aplikasi dan versi yang kamu pakai, "
               "serta cara membaca hasilnya, gunakan skill panduan-analisis. Pilihan teknik tetap ditentukan dari rancangan penelitianmu "
               "bersama pembimbing.", 10, italic=True)


def buat_docx(kode, f, keluar, baris=None, skala=None):
    doc = _new_doc(f["landscape"])
    _para(doc, f["judul"].upper(), 14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(doc, "Kategori penggunaan " + f["kategori"], 10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _identitas_table(doc, f.get("identitas", []))
    kolom = list(f["kolom"])
    n = baris if baris is not None else f["baris"]
    if kode == "angket-kosong":
        s = skala or 5
        kolom = ["No", "Pernyataan"] + [str(i) for i in range(1, s + 1)]
    if "bagian" in f:
        _main_table(doc, kolom, 0, bagian=f["bagian"])
    else:
        _main_table(doc, kolom, n, total_cm=25.7 if f["landscape"] else 17.0)
    _panduan_pages(doc, f)
    out = Path(keluar) / f"belangko-{kode}.docx"
    doc.save(out)
    return out


# ---------- util xlsx ----------
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _xl_header(ws, row, cols):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(name="Arial", bold=True, size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(c) + 6))


def _xl_empty_grid(ws, first_row, n_rows, n_cols, fixed_first=None):
    for i in range(n_rows):
        for j in range(1, n_cols + 1):
            cell = ws.cell(row=first_row + i, column=j)
            cell.border = BORDER
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fixed_first:
            ws.cell(row=first_row + i, column=1, value=fixed_first(i))


def _xl_petunjuk(wb, f, judul):
    ws = wb.active
    ws.title = "Petunjuk"
    ws["A1"] = judul
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    r = 3
    def blok(titel, items, num):
        nonlocal r
        ws.cell(row=r, column=1, value=titel).font = Font(name="Arial", bold=True, size=11)
        r += 1
        for n, it in enumerate(items, start=1):
            c = ws.cell(row=r, column=1, value=(f"{n}. " if num else "- ") + it)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name="Arial", size=10)
            r += 1
        r += 1
    blok("Petunjuk Pengisian", f["petunjuk"] + [ETIKA_DATA], True)
    blok("Bahan yang Dibutuhkan untuk Analisis", f["bahan"], False)
    blok("Langkah Analisis Umum", f["langkah"], True)
    c = ws.cell(row=r, column=1, value="Catatan. Untuk tutorial menurut aplikasi dan versi yang kamu pakai, gunakan skill panduan-analisis. " + CATATAN_AKHIR)
    c.font = Font(name="Arial", italic=True, size=9); c.alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 110
    return ws


def buat_xlsx(kode, f, keluar, butir=None, responden=None, baris=None):
    wb = Workbook()
    _xl_petunjuk(wb, f, f["judul"])
    if kode == "tabulasi-angket":
        nb, nr = butir or 20, responden or 30
        ws = wb.create_sheet("Kamus Variabel")
        _xl_header(ws, 1, XLSX_KOLOM["kamus"])
        _xl_empty_grid(ws, 2, nb, len(XLSX_KOLOM["kamus"]), lambda i: f"B{i+1:02d}")
        ws.column_dimensions["B"].width = 50
        ws2 = wb.create_sheet("Data")
        cols = ["No", "Kode Responden"] + [f"B{i:02d}" for i in range(1, nb + 1)]
        _xl_header(ws2, 1, cols)
        for j in range(3, len(cols) + 1):
            ws2.column_dimensions[get_column_letter(j)].width = 7
        _xl_empty_grid(ws2, 2, nr, len(cols), lambda i: i + 1)
        ws2.freeze_panes = "C2"
    elif kode == "tabulasi-tes":
        nb, nr = butir or 20, responden or 30
        ws = wb.create_sheet("Kunci")
        _xl_header(ws, 1, XLSX_KOLOM["kunci"])
        _xl_empty_grid(ws, 2, nb, len(XLSX_KOLOM["kunci"]), lambda i: i + 1)
        ws2 = wb.create_sheet("Data Jawaban")
        cols = ["No", "Kode Peserta"] + [f"S{i:02d}" for i in range(1, nb + 1)]
        _xl_header(ws2, 1, cols)
        for j in range(3, len(cols) + 1):
            ws2.column_dimensions[get_column_letter(j)].width = 7
        _xl_empty_grid(ws2, 2, nr, len(cols), lambda i: i + 1)
        ws2.freeze_panes = "C2"
    elif kode == "tabulasi-pretest-posttest":
        nr = responden or 30
        ws = wb.create_sheet("Parameter")
        _xl_header(ws, 1, ["Parameter", "Nilai"])
        for i, lab in enumerate(["Skor maksimum tes", "Tanggal pretest", "Tanggal posttest", "Keterangan kelompok"], start=2):
            ws.cell(row=i, column=1, value=lab).font = Font(name="Arial", bold=True, size=10)
            ws.cell(row=i, column=2).border = BORDER
            ws.cell(row=i, column=1).border = BORDER
        ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 30
        ws2 = wb.create_sheet("Data")
        _xl_header(ws2, 1, XLSX_KOLOM["prepost"])
        _xl_empty_grid(ws2, 2, nr, len(XLSX_KOLOM["prepost"]), lambda i: i + 1)
        ws2.freeze_panes = "A2"
    elif kode == "matriks-literatur":
        n = baris or 25
        ws = wb.create_sheet("Matriks")
        _xl_header(ws, 1, XLSX_KOLOM["literatur"])
        _xl_empty_grid(ws, 2, n, len(XLSX_KOLOM["literatur"]), lambda i: i + 1)
        for col, w in zip("BCDEFGHIJK", (24, 36, 24, 26, 30, 24, 40, 30, 30, 18)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "C2"
    out = Path(keluar) / f"belangko-{kode}.xlsx"
    wb.save(out)
    return out


# ---------- uji kekosongan ----------
def uji_docx(path, kode):
    f = FORMS[kode]
    d = Document(path)
    masalah = []
    tabel = list(d.tables)
    idx = 0
    if f.get("identitas"):
        t = tabel[idx]; idx += 1
        for r in t.rows:
            if r.cells[1].text.strip():
                masalah.append(f"identitas terisi: {r.cells[0].text}")
    t = tabel[idx]
    for ri, r in enumerate(t.rows):
        if ri == 0:
            continue
        for ci, c in enumerate(r.cells):
            txt = c.text.strip()
            if "bagian" in f:
                if ci == 1 and txt:
                    masalah.append(f"sel terisi baris {ri}")
            else:
                no_col = (ci == 0 and f["kolom"][0].lower().startswith("no"))
                if no_col:
                    if txt and not txt.isdigit():
                        masalah.append(f"kolom No berisi teks baris {ri}")
                elif txt:
                    masalah.append(f"sel terisi baris {ri} kolom {ci}: {txt[:30]}")
    return masalah


def uji_xlsx(path, kode):
    wb = load_workbook(path)
    masalah = []
    for ws in wb.worksheets:
        if ws.title == "Petunjuk":
            continue
        hdr = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value is None:
                    continue
                v = str(c.value)
                kolom_awal = c.column == 1
                if ws.title == "Parameter":
                    if kolom_awal: continue
                    masalah.append(f"{ws.title}!{c.coordinate} terisi")
                    continue
                if kolom_awal and (v.isdigit() or re.fullmatch(r"B\d{2}", v)):
                    continue
                masalah.append(f"{ws.title}!{c.coordinate} terisi: {v[:30]}")
        # larang rumus
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    masalah.append(f"{ws.title}!{c.coordinate} berisi rumus")
    return masalah


# ---------- katalog markdown ----------
def katalog_md():
    out = ["# Katalog Belangko Penelitian", "",
           "Dibuat otomatis dari `scripts/katalog.py`. Semua belangko kosong, hanya format.", ""]
    for kode, f in FORMS.items():
        out += [f"## {f['judul']}", "",
                f"Kode `{kode}`. Format {f['jenis']}. Kategori {f['kategori']}.", "",
                "Bahan yang dibutuhkan untuk analisis.", ""]
        out += [f"- {b}" for b in f["bahan"]] + ["", "Langkah analisis umum.", ""]
        out += [f"{i}. {s}" for i, s in enumerate(f["langkah"], start=1)] + [""]
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--daftar", action="store_true")
    ap.add_argument("--jenis")
    ap.add_argument("--semua", action="store_true")
    ap.add_argument("--keluar", default=".")
    ap.add_argument("--butir", type=int)
    ap.add_argument("--responden", type=int)
    ap.add_argument("--baris", type=int)
    ap.add_argument("--skala", type=int)
    ap.add_argument("--uji", action="store_true")
    ap.add_argument("--katalog-md", action="store_true")
    a = ap.parse_args()

    if a.daftar:
        for k, f in FORMS.items():
            print(f"{k:32s} {f['jenis']:5s} {f['judul']}")
        return
    if a.katalog_md:
        print(katalog_md()); return
    os.makedirs(a.keluar, exist_ok=True)
    kodes = list(FORMS) if a.semua else ([a.jenis] if a.jenis else [])
    if a.uji and not kodes:
        kodes = list(FORMS)
    if not kodes:
        ap.error("Pilih --jenis KODE, --semua, atau --daftar")
    gagal = 0
    for k in kodes:
        if k not in FORMS:
            sys.exit(f"Kode tidak dikenal. Lihat --daftar. {k}")
        f = FORMS[k]
        if a.uji:
            p = Path(a.keluar) / (f"belangko-{k}." + f["jenis"])
            if not p.exists():
                p = (buat_docx if f["jenis"] == "docx" else buat_xlsx)(k, f, a.keluar)
            m = uji_docx(p, k) if f["jenis"] == "docx" else uji_xlsx(p, k)
            print(("LULUS " if not m else "GAGAL ") + p.name + ("" if not m else " " + "; ".join(m[:5])))
            gagal += bool(m)
            continue
        if f["jenis"] == "docx":
            p = buat_docx(k, f, a.keluar, baris=a.baris, skala=a.skala)
        else:
            p = buat_xlsx(k, f, a.keluar, butir=a.butir, responden=a.responden, baris=a.baris)
        print("dibuat", p)
    if a.uji:
        sys.exit(1 if gagal else 0)


if __name__ == "__main__":
    main()
